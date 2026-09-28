"""Test helpers: tiny in-memory document builders and an SSE parser."""

from __future__ import annotations

import io
import json
from typing import Any

import docx


class CharTokenizer:
    """One token per character: makes token budgets in chunking tests easy to reason about."""

    def encode(self, text: str) -> list[int]:
        return [ord(ch) for ch in text]

    def decode(self, tokens: list[int]) -> str:
        return "".join(chr(t) for t in tokens)


def make_pdf(pages: list[str], title: str | None = None) -> bytes:
    """Build a minimal, valid PDF with one line of Helvetica text per page."""
    objects: list[bytes] = []
    n_pages = len(pages)
    font_id = 3 + 2 * n_pages
    info_id = font_id + 1
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(f"<< /Title ({title or ''}) >>".encode())

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R"
    if title:
        trailer += f" /Info {info_id} 0 R"
    out.write(f"{trailer} >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(title: str, sections: dict[str, list[str]]) -> bytes:
    document = docx.Document()
    document.add_heading(title, level=0)
    for heading, paragraphs in sections.items():
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse_sse(body: str) -> list[tuple[str, Any]]:
    """Parse a text/event-stream body into (event, json_data) pairs, skipping comments."""
    events: list[tuple[str, Any]] = []
    for frame in body.replace("\r\n", "\n").split("\n\n"):
        event, data_lines = "message", []
        for line in frame.split("\n"):
            if line.startswith("event:"):
                event = line[6:].strip()
            elif line.startswith("data:"):
                data_lines.append(line[5:].lstrip())
        if data_lines:
            events.append((event, json.loads("\n".join(data_lines))))
    return events
