"""Extract text blocks (with page / heading metadata) from supported file types."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.ingestion.cleaning import clean_text

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
}

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_MD_FRONT_MATTER = re.compile(r"\A---\n.*?\n---\n", re.DOTALL)


class DocumentParseError(ValueError):
    """The file is corrupt, unsupported, or contains no extractable text."""


@dataclass(frozen=True, slots=True)
class TextBlock:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    title: str
    blocks: list[TextBlock] = field(default_factory=list)
    page_count: int | None = None


def detect_file_type(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    try:
        return SUPPORTED_EXTENSIONS[suffix]
    except KeyError:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentParseError(
            f"Unsupported file type '{suffix or filename}'. Supported: {supported}"
        ) from None


def _default_title(filename: str) -> str:
    stem = PurePath(filename).stem
    return re.sub(r"[_\-]+", " ", stem).strip() or filename


def parse_document(data: bytes, filename: str) -> ParsedDocument:
    file_type = detect_file_type(filename)
    parser = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "markdown": _parse_markdown,
        "text": _parse_text,
    }[file_type]
    parsed = parser(data, filename)
    parsed.blocks = [
        TextBlock(text=cleaned, page=block.page, heading=block.heading)
        for block in parsed.blocks
        if (cleaned := clean_text(block.text))
    ]
    if not parsed.blocks:
        raise DocumentParseError("No extractable text found in the document")
    return parsed


def _parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        blocks = [
            TextBlock(text=page.extract_text() or "", page=number)
            for number, page in enumerate(reader.pages, start=1)
        ]
        meta_title = reader.metadata.title if reader.metadata else None
    except PdfReadError as exc:
        raise DocumentParseError(f"Could not read PDF: {exc}") from exc
    if not any(block.text.strip() for block in blocks):
        raise DocumentParseError("The PDF has no text layer (is it a scan?). OCR is not supported.")
    title = meta_title.strip() if meta_title and meta_title.strip() else _default_title(filename)
    return ParsedDocument(title=title, blocks=blocks, page_count=len(blocks))


def _docx_items(document: DocxDocument) -> list[Paragraph | Table]:
    """Paragraphs and tables in body order."""
    items: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            items.append(Paragraph(child, document))
        elif tag == "tbl":
            items.append(Table(child, document))
    return items


def _parse_docx(data: bytes, filename: str) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of zip/xml errors
        raise DocumentParseError(f"Could not read DOCX: {exc}") from exc

    title = (document.core_properties.title or "").strip()
    blocks: list[TextBlock] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append(TextBlock(text="\n\n".join(buffer), heading=heading))
            buffer.clear()

    for item in _docx_items(document):
        if isinstance(item, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in item.rows]
            buffer.append("\n".join(rows))
            continue
        text = item.text.strip()
        if not text:
            continue
        style = item.style.name if item.style is not None and item.style.name else ""
        if style == "Title":
            title = title or text
        elif style.startswith("Heading"):
            flush()
            heading = text
        else:
            buffer.append(text)
    flush()
    return ParsedDocument(title=title or _default_title(filename), blocks=blocks)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _parse_markdown(data: bytes, filename: str) -> ParsedDocument:
    text = _MD_FRONT_MATTER.sub("", _decode_text(data).replace("\r\n", "\n"))
    title: str | None = None
    stack: list[tuple[int, str]] = []
    blocks: list[TextBlock] = []
    buffer: list[str] = []
    in_fence = False

    def current_heading() -> str | None:
        # The document title (first H1) is already shown alongside every chunk; omit it.
        path = [name for level, name in stack if not (level == 1 and name == title)]
        return " > ".join(path) or None

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            blocks.append(TextBlock(text=body, heading=current_heading()))
        buffer.clear()

    for line in text.split("\n"):
        if _MD_FENCE.match(line):
            in_fence = not in_fence
            buffer.append(line)
            continue
        match = None if in_fence else _MD_HEADING.match(line)
        if match is None:
            buffer.append(line)
            continue
        flush()
        level, name = len(match.group(1)), match.group(2).strip()
        if level == 1 and title is None:
            title = name
        while stack and stack[-1][0] >= level:
            stack.pop()
        stack.append((level, name))
    flush()
    return ParsedDocument(title=title or _default_title(filename), blocks=blocks)


def _parse_text(data: bytes, filename: str) -> ParsedDocument:
    return ParsedDocument(
        title=_default_title(filename), blocks=[TextBlock(text=_decode_text(data))]
    )
